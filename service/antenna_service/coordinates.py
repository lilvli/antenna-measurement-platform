from __future__ import annotations

import hashlib
import json
import math
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from antenna_service.errors import ServiceError, invalid


@dataclass(frozen=True, slots=True)
class Channel:
    polarization: str
    element: int
    spi_no: int
    chip_no: int
    chip_channel_index: int
    grid_row: int
    grid_column: int
    x: float
    y: float
    z: float
    enabled: bool


@dataclass(slots=True)
class CoordinateModel:
    asset_id: str
    path: str
    file_sha256: str
    schema_version: str
    release_status: str
    antenna_id: str
    unit: str
    evidence_limit: str
    evidence_description: str
    channels: list[Channel]
    mapping_sha256: str
    enabled_sha256: str
    geometry_sha256: str

    @property
    def polarizations(self) -> list[str]:
        return sorted({channel.polarization for channel in self.channels})

    @property
    def evidence(self) -> str:
        text = f"{self.release_status} {self.evidence_limit} {self.evidence_description}".upper()
        return "SIMULATED" if "SYNTHETIC" in text or "SIMULATED" in text or "NOT_VERIFIED" in text else "REAL"

    def for_polarization(self, polarization: str) -> list[Channel]:
        return sorted(
            [channel for channel in self.channels if channel.polarization == polarization.upper()],
            key=lambda channel: channel.element,
        )

    def summary(self) -> dict[str, Any]:
        # The renderer only needs the grid dimensions for each polarization.
        # Keeping this small summary avoids sending all channel coordinates again.
        polarization_layouts: dict[str, dict[str, int]] = {}
        for polarization in self.polarizations:
            channels = self.for_polarization(polarization)
            polarization_layouts[polarization] = {
                "channel_count": len(channels),
                "enabled_count": sum(channel.enabled for channel in channels),
                "rows": len({channel.grid_row for channel in channels}),
                "columns": len({channel.grid_column for channel in channels}),
            }
        return {
            "asset_id": self.asset_id,
            "antenna_id": self.antenna_id,
            "schema_version": self.schema_version,
            "release_status": self.release_status,
            "evidence": self.evidence,
            "evidence_limit": self.evidence_limit,
            "unit": self.unit,
            "polarizations": self.polarizations,
            "channel_count": len(self.channels),
            "enabled_count": sum(channel.enabled for channel in self.channels),
            "polarization_layouts": polarization_layouts,
            "file_sha256": self.file_sha256,
            "mapping_sha256": self.mapping_sha256,
            "enabled_sha256": self.enabled_sha256,
            "geometry_sha256": self.geometry_sha256,
        }


def _hash_json(value: Any) -> str:
    encoded = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _file_sha(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class CoordinateLoader:
    EXPECTED_HEADER = [
        "极化",
        "天线单元",
        "SPI号",
        "芯片号",
        "芯片通道索引（0基）",
        "标准行",
        "标准列",
        "X",
        "Y",
        "Z",
        "通道使能",
    ]

    def load(self, file_path: str) -> CoordinateModel:
        path = Path(file_path).resolve()
        if path.suffix.lower() != ".xlsx":
            raise invalid("坐标表必须是 .xlsx", stage="coordinates_load", target=str(path))
        try:
            workbook = load_workbook(path, read_only=False, data_only=True)
        except Exception as exc:
            raise ServiceError("DATA_INTEGRITY", "无法读取坐标表", "coordinates_load", str(path), {"error": str(exc)}) from exc
        if workbook.sheetnames != ["通道坐标表"]:
            raise invalid(
                "坐标工作簿必须且只能包含“通道坐标表”工作表",
                stage="coordinates_load",
                target=str(path),
                sheets=workbook.sheetnames,
            )
        sheet = workbook["通道坐标表"]
        if any(state != "visible" for state in [sheet.sheet_state]):
            raise invalid("业务工作表不能隐藏", stage="coordinates_load", target=str(path))
        metadata = {
            str(sheet.cell(4, column).value).strip(): sheet.cell(4, column + 1).value
            for column in range(1, 10, 2)
        }
        schema = str(metadata.get("格式版本", "")).strip()
        if schema not in {"2.1.0", "2.0.0", "2.0"}:
            raise invalid("不支持的坐标表格式版本", stage="coordinates_load", target="通道坐标表!B4", value=schema)
        header = [sheet.cell(8, column).value for column in range(1, 12)]
        if schema == "2.1.0" and header != self.EXPECTED_HEADER:
            raise invalid(
                "2.1.0 坐标表表头损坏或缺列",
                stage="coordinates_load",
                target="通道坐标表!A8:K8",
                expected=self.EXPECTED_HEADER,
                actual=header,
            )
        legacy = schema != "2.1.0"
        required = self.EXPECTED_HEADER[:-1] if legacy else self.EXPECTED_HEADER
        if header[: len(required)] != required:
            raise invalid("坐标表表头不合法", stage="coordinates_load", target="通道坐标表!A8")

        channels: list[Channel] = []
        for row in range(9, sheet.max_row + 1):
            values = [sheet.cell(row, column).value for column in range(1, 12)]
            if all(value in (None, "") for value in values):
                continue
            try:
                polarization = str(values[0]).strip().upper()
                integer_values = []
                for index in range(1, 7):
                    raw = values[index]
                    if isinstance(raw, bool) or not isinstance(raw, (int, float)) or not math.isfinite(raw) or int(raw) != raw:
                        raise ValueError(f"{self.EXPECTED_HEADER[index]}必须是整数")
                    integer_values.append(int(raw))
                coordinates = [float(values[index]) for index in range(7, 10)]
                enabled_raw = 1 if legacy else values[10]
                if isinstance(enabled_raw, bool) or enabled_raw not in (0, 1):
                    raise ValueError("通道使能必须是严格整数 0 或 1")
                if polarization not in {"H", "V"}:
                    raise ValueError("极化必须是 H 或 V")
                if any(value < 0 for value in integer_values) or not all(math.isfinite(value) for value in coordinates):
                    raise ValueError("编号必须非负且坐标必须为有限数")
            except (TypeError, ValueError) as exc:
                raise invalid(
                    f"坐标行数据非法：{exc}",
                    stage="coordinates_load",
                    target=f"通道坐标表!A{row}:K{row}",
                ) from exc
            channels.append(
                Channel(
                    polarization,
                    integer_values[0],
                    integer_values[1],
                    integer_values[2],
                    integer_values[3],
                    integer_values[4],
                    integer_values[5],
                    coordinates[0],
                    coordinates[1],
                    coordinates[2],
                    bool(enabled_raw),
                )
            )
        if not channels:
            raise invalid("坐标表没有通道数据", stage="coordinates_load", target="通道坐标表")
        self._validate(channels)

        file_sha = _file_sha(path)
        mapping = [
            (c.polarization, c.element, c.spi_no, c.chip_no, c.chip_channel_index, c.grid_row, c.grid_column)
            for c in channels
        ]
        enabled = [(c.polarization, c.element, int(c.enabled)) for c in channels]
        geometry = [(c.polarization, c.element, c.x, c.y, c.z) for c in channels]
        antenna_id = str(metadata.get("天线ID", "")).strip()
        if not antenna_id:
            raise invalid("天线ID不能为空", stage="coordinates_load", target="通道坐标表!F4")
        return CoordinateModel(
            asset_id=f"{antenna_id}:{file_sha[:12]}",
            path=str(path),
            file_sha256=file_sha,
            schema_version=schema,
            release_status=str(metadata.get("发布状态", "")).strip(),
            antenna_id=antenna_id,
            unit=str(metadata.get("坐标单位", "")).strip(),
            evidence_limit=str(metadata.get("证据上限", "")).strip(),
            evidence_description=str(sheet.cell(6, 2).value or "").strip(),
            channels=channels,
            mapping_sha256=_hash_json(mapping),
            enabled_sha256=_hash_json(enabled),
            geometry_sha256=_hash_json(geometry),
        )

    def _validate(self, channels: list[Channel]) -> None:
        mapping_keys: set[tuple[int, int, int]] = set()
        identities: set[tuple[str, int]] = set()
        grid_keys: set[tuple[str, int, int]] = set()
        for channel in channels:
            mapping = (channel.spi_no, channel.chip_no, channel.chip_channel_index)
            identity = (channel.polarization, channel.element)
            grid = (channel.polarization, channel.grid_row, channel.grid_column)
            if mapping in mapping_keys:
                raise invalid("SPI/芯片/芯片通道三元组重复", stage="coordinates_validate", target=str(mapping))
            if identity in identities:
                raise invalid("极化+天线单元重复", stage="coordinates_validate", target=str(identity))
            if grid in grid_keys:
                raise invalid("极化+标准行列重复", stage="coordinates_validate", target=str(grid))
            mapping_keys.add(mapping)
            identities.add(identity)
            grid_keys.add(grid)
        for polarization in sorted({channel.polarization for channel in channels}):
            subset = [channel for channel in channels if channel.polarization == polarization]
            elements = sorted(channel.element for channel in subset)
            if elements != list(range(len(elements))):
                raise invalid(
                    "每个极化的天线单元必须从 0 连续编号",
                    stage="coordinates_validate",
                    target=polarization,
                )
            rows = sorted({channel.grid_row for channel in subset})
            columns = sorted({channel.grid_column for channel in subset})
            expected = {(polarization, row, column) for row in rows for column in columns}
            actual = {(polarization, channel.grid_row, channel.grid_column) for channel in subset}
            if actual != expected:
                raise invalid("标准行列必须形成完整矩形网格", stage="coordinates_validate", target=polarization)


def validate_profile_coordinate_pair(profile: Any, coordinates: CoordinateModel) -> None:
    supported = profile.supported_polarizations
    actual = set(coordinates.polarizations)
    if supported in {"SINGLE", "SINGLE_POLARIZATION"}:
        if len(actual) != 1:
            raise ServiceError(
                "NOT_RUNNABLE", "单极化配置包不能配合双极化坐标表", "asset_pair", coordinates.antenna_id
            )
    else:
        declared = {part.strip() for part in supported.replace("+", ",").split(",") if part.strip()}
        if not actual.issubset(declared):
            raise ServiceError(
                "NOT_RUNNABLE",
                "坐标表极化与配置包声明不兼容",
                "asset_pair",
                coordinates.antenna_id,
                {"profile": sorted(declared), "coordinates": sorted(actual)},
            )
    if profile.capabilities.get("calibration"):
        # Compile every enabled mapping for both paths before a real device can be touched.
        for channel in coordinates.channels:
            if not channel.enabled:
                continue
            for signal_path in ("TX", "RX"):
                profile.build_calibration_frame(
                    array_id=0,
                    spi_no=channel.spi_no,
                    chip_no=channel.chip_no,
                    chip_channel_index=channel.chip_channel_index,
                    signal_path=signal_path,
                )
