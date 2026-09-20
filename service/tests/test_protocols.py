from __future__ import annotations

import json
import queue
import threading
from collections import deque

import pytest
from openpyxl import load_workbook

from antenna_service.coordinates import validate_profile_coordinate_pair
from antenna_service.errors import ServiceError
from antenna_service.protocol.profile import ProfileLoader
from antenna_service.protocol.crc import append_crc_be, crc16_ccitt_false, verify_crc_be
from antenna_service.protocol.flash import (
    PAGE_READ_OPCODE,
    PAGE_WRITE_OPCODE,
    decode_page_response,
    encode_page_read,
    encode_page_write,
)
from antenna_service.protocol.rtc import OPCODES, RtcClient, RtcEndpoint, build_fixed, parse_frame


def _find_row(sheet, column: int, value: str) -> int:
    for row in range(1, sheet.max_row + 1):
        if str(sheet.cell(row, column).value or "").strip() == value:
            return row
    raise AssertionError(f"未找到 {sheet.title}!{column} 列中的 {value}")


def test_profile_vectors_and_all_coordinate_mappings(loaded_assets):
    _, profile, coordinates = loaded_assets
    validate_profile_coordinate_pair(profile, coordinates)
    assert all(result["passed"] for result in profile.vector_results)
    assert len(coordinates.channels) == 256
    summary = coordinates.summary()
    assert summary["enabled_count"] == 256
    assert summary["polarization_layouts"] == {
        "H": {"channel_count": 256, "enabled_count": 256, "rows": 16, "columns": 16}
    }
    assert profile.build_calibration_frame(
        array_id=0,
        spi_no=3,
        chip_no=2,
        chip_channel_index=0,
        signal_path="TX",
    ) == bytes.fromhex("AA 55 00 16 31 00 00 03 22 42 01 FF 00 00 00 00 00 00 00 00 EA 95")
    profile_summary = profile.summary()
    assert profile_summary["commands"]
    assert all("fields" in command and "opcode_hex" in command for command in profile_summary["commands"])
    json.dumps(profile_summary, ensure_ascii=False)


def test_pattern_signal_path_is_encoded_as_rx_00_or_tx_ff(loaded_assets):
    _, profile, _ = loaded_assets
    command = profile.command_for_role("BEAM_SET")
    inputs = {
        "off_axis_deg": 45,
        "azimuth_deg": 30.1,
        "frequency_ghz": 1,
    }
    rx = profile.encode(command.command_id, 0, inputs | {"signal_path": "RX"})
    tx = profile.encode(command.command_id, 0, inputs | {"signal_path": "TX"})

    assert len(rx) == len(tx) == 22
    assert rx[13] == 0x00
    assert tx[13] == 0xFF
    assert rx[14:17] == rx[17:20] == bytes.fromhex("00 04 00")
    assert tx[14:17] == tx[17:20] == bytes.fromhex("00 04 00")
    assert tx == bytes.fromhex("AA 55 00 16 22 00 00 00 11 94 0B C2 00 FF 00 04 00 00 04 00 05 ED")


def test_pattern_signal_path_can_be_omitted_and_unoccupied_byte_defaults_to_zero(loaded_assets, tmp_path):
    _, loaded_profile, _ = loaded_assets
    profile_path = tmp_path / "optional_signal_path.xlsx"
    workbook = load_workbook(loaded_profile.path)
    send_sheet = workbook["03_发送字段"]
    beam_command = loaded_profile.command_for_role("BEAM_SET").command_id
    signal_row = next(
        row
        for row in range(1, send_sheet.max_row + 1)
        if str(send_sheet.cell(row, 2).value or "").strip() == beam_command
        and str(send_sheet.cell(row, 14).value or "").strip() == "signal_path"
    )
    send_sheet.cell(signal_row, 1).value = "N"
    vector_sheet = workbook["06_测试向量"]
    for row in range(1, vector_sheet.max_row + 1):
        if str(vector_sheet.cell(row, 4).value or "").strip() == beam_command:
            vector_sheet.cell(row, 1).value = "N"
    workbook.save(profile_path)

    profile = ProfileLoader().load(str(profile_path))
    assert profile.capabilities["pattern"] is True
    assert profile.capabilities["beam_signal_path"] is False
    frame = profile.encode(
        beam_command,
        0,
        {
            "off_axis_deg": 45,
            "azimuth_deg": 30.1,
            "frequency_ghz": 1,
            "signal_path": "TX",
        },
    )
    assert len(frame) == 22
    assert frame[13] == 0x00


def test_receive_sections_are_located_by_markers_instead_of_fixed_rows(loaded_assets, tmp_path):
    _, loaded_profile, _ = loaded_assets
    profile_path = tmp_path / "compact_receive_sections.xlsx"
    workbook = load_workbook(loaded_profile.path)
    sheet = workbook["04_接收解析"]
    for merged_range in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged_range))
    fields_marker_row = _find_row(sheet, 1, "B. 应答字段")
    last_rule_row = max(
        row
        for row in range(1, fields_marker_row)
        if str(sheet.cell(row, 1).value or "").strip().upper() == "Y"
    )
    sheet.delete_rows(last_rule_row + 1, fields_marker_row - last_rule_row - 1)
    workbook.save(profile_path)

    profile = ProfileLoader().load(str(profile_path))
    assert profile.response_rules
    assert profile.response_fields["RULE_BEAM_CONTROL"]


def test_receive_section_marker_error_identifies_sheet_and_cell(loaded_assets, tmp_path):
    _, loaded_profile, _ = loaded_assets
    profile_path = tmp_path / "missing_receive_marker.xlsx"
    workbook = load_workbook(loaded_profile.path)
    sheet = workbook["04_接收解析"]
    marker_row = _find_row(sheet, 1, "B. 应答字段")
    sheet.cell(marker_row, 1).value = "应答字段"
    workbook.save(profile_path)

    with pytest.raises(ServiceError) as captured:
        ProfileLoader().load(str(profile_path))
    assert captured.value.target == "04_接收解析!A:A"


def test_old_profile_workbook_is_rejected_without_compatibility_layer(loaded_assets, tmp_path):
    _, loaded_profile, _ = loaded_assets
    old_profile = tmp_path / "old_profile.xlsx"
    workbook = load_workbook(loaded_profile.path)
    workbook["01_基本信息"]["B5"] = "V0.9"
    workbook.save(old_profile)
    with pytest.raises(ServiceError, match="V1.0"):
        ProfileLoader().load(str(old_profile))


def test_crc_and_rtc_fixed_vectors():
    assert crc16_ccitt_false(b"123456789") == 0x29B1
    wave = bytes.fromhex("AA 55 00 16 22 00 00 00 11 94 0B B8 00 00 00 04 00 00 04 00 E4 DF")
    assert build_fixed(OPCODES["WRITE_WAVE_ENTRY"], aux=1, data=wave) == bytes.fromhex(
        "A5 5A 00 20 21 00 00 01 AA 55 00 16 22 00 00 00 11 94 0B B8 00 00 00 04 00 00 04 00 E4 DF D8 E1"
    )
    transfer = build_fixed(OPCODES["ANTENNA_TRANSFER_22"], data=wave)
    assert len(transfer) == 32
    assert transfer[-2:] == bytes.fromhex("76 D1")
    assert verify_crc_be(transfer)


def test_flash_page_frames_and_correlated_responses():
    tile_id = 7
    address = 0x12AB00
    page = bytes(index % 256 for index in range(256))
    write_request = encode_page_write(tile_id, address, page)
    assert len(write_request) == 286
    assert write_request[:11] == bytes.fromhex("AA 55 01 1E 0A 07 00 32 12 AB 00")
    assert write_request[11:267] == page
    assert write_request[267:284] == bytes(17)
    assert verify_crc_be(write_request)

    write_ack = append_crc_be(bytes.fromhex("AA 55 00 16 0A 07 00 32 12 AB 00") + bytes(9))
    assert decode_page_response(write_ack, PAGE_WRITE_OPCODE, tile_id, address) == b""

    read_request = encode_page_read(tile_id, address)
    assert len(read_request) == 22
    assert read_request[:11] == bytes.fromhex("AA 55 00 16 0A 07 00 6B 12 AB 00")
    read_response = append_crc_be(bytes.fromhex("AA 55 01 1E 0A 07 00 6B 12 AB 00") + page + bytes(17))
    assert decode_page_response(read_response, PAGE_READ_OPCODE, tile_id, address) == page


def test_receive_parser_only_matches_loaded_profile_and_returns_compact_fields(loaded_assets):
    registry, _, _ = loaded_assets
    frame = bytes.fromhex("AA 55 00 16 22 00 00 00 11 94 0B B8 00 00 00 04 00 00 04 00 E4 DF")
    parsed = registry.decode_matching_frame(frame)
    assert parsed is not None
    assert parsed["command_id"]
    assert parsed["fields"]
    assert all(set(field) == {"key", "label", "value", "unit"} for field in parsed["fields"])

    unknown_prefix = bytearray(frame[:-2])
    unknown_prefix[4] = 0xFE
    assert registry.decode_matching_frame(append_crc_be(bytes(unknown_prefix))) is None

    damaged = bytearray(frame)
    damaged[-1] ^= 0x01
    assert registry.decode_matching_frame(bytes(damaged)) is None


class FakeRtcSerial:
    def __init__(self, frames):
        self.frames = deque(frames)
        self.incoming = queue.Queue()
        self.sent = []

    def write(self, frame):
        self.sent.append(frame)
        while self.frames:
            self.incoming.put(self.frames.popleft())
        return len(frame)

    def read(self, length):
        try:
            return self.incoming.get(timeout=0.01)
        except queue.Empty:
            return b""

    def close(self):
        pass


def attach_fake_serial(client, frames):
    client._serial = fake = FakeRtcSerial(frames)
    client._reader = threading.Thread(target=client._receive_loop, daemon=True)
    client._reader.start()
    return fake



@pytest.mark.asyncio
async def test_rtc_dispatches_events_and_unrelated_ack_without_resending_query():
    client = RtcClient(RtcEndpoint(port="FAKE", timeout_ms=50))
    event = build_fixed(0xE2, control=22, data=bytes(22))
    unrelated = build_fixed(0x81, data=bytes(22))
    status = build_fixed(0x83, data=bytes([1]) + bytes(21))
    fake = attach_fake_serial(client, [event, unrelated, status])
    try:
        result = await client.get_status()
        assert result["state"] == "IDLE"
        assert result["events"][0]["event"] == "ANTENNA_RX"
        assert result["events"][0]["raw_hex"] == event.hex(" ").upper()
        assert len(fake.sent) == 1
        json.dumps(result)
    finally:
        await client.close()



@pytest.mark.asyncio
async def test_rtc_ignores_stale_address_and_unrelated_nack():
    client = RtcClient(RtcEndpoint(port="FAKE", timeout_ms=50))
    wave = bytes(range(22))
    wrong_address = build_fixed(0xA4, aux=1, data=bytes(22))
    wrong_nack = build_fixed(0x7F, control=6, aux=OPCODES["PING"] << 8)
    correct = build_fixed(0xA4, aux=2, data=wave)
    fake = attach_fake_serial(client, [wrong_address, wrong_nack, correct])
    try:
        assert await client.read_wave_entry(2) == wave
        assert len(fake.sent) == 1
    finally:
        await client.close()



@pytest.mark.asyncio
async def test_rtc_transfer_crc_is_diagnostic_and_never_retries_on_timeout():
    client = RtcClient(RtcEndpoint(port="FAKE", timeout_ms=50))
    request = build_fixed(OPCODES["ANTENNA_TRANSFER_22"], data=bytes(22))
    diagnostic_crc = b"\x00\x16\x12\x34" + bytes(18)
    fake = attach_fake_serial(client, [build_fixed(0xB0, data=diagnostic_crc)])
    try:
        response = await client.transact(request)
        assert response[10:12] == b"\x12\x34"
        assert len(fake.sent) == 1
        with pytest.raises(ServiceError) as error:
            await client.transact(request)
        assert error.value.code == "TIMEOUT"
        assert error.value.side_effect_possible
        assert len(fake.sent) == 2
    finally:
        await client.close()



def test_rtc_rejects_short_crc_valid_fixed_responses():
    short = append_crc_be(bytes.fromhex("A5 5A 00 0A 83 00 00 00"))
    with pytest.raises(ServiceError, match="长度错误"):
        parse_frame(short)
