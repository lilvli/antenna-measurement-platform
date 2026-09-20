from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from antenna_service.errors import ServiceError
from antenna_service.protocol.profile import ProfileLoader
from antenna_service.protocol.profile import AssetRegistry
from conftest import PROFILE_PATH


def row(sheet, index, values):
    for col,value in enumerate(values,1):
        sheet.cell(index,col).value=value


def command(sheet, index, name='EXTRA', opcode='70'):
    row(sheet,index,['Y',name,'扩展命令','MANUAL_ONLY',opcode,'SINGLE',opcode,1000,'SAFE','VALID_RESPONSE',None,None,'N','扩展测试'])


def field(sheet,index,name='EXTRA',kind='uint8',length=1,default=77,enum=None,layout=None):
    row(sheet,index,['Y',name,'extra','扩展字段',1,length,kind,1,None,0,255,'USER',default,None,enum,'扩展测试',layout])


def save_load(workbook,tmp_path):
    path=tmp_path/'profile.xlsx'
    workbook.save(path)
    return ProfileLoader().load(str(path))


def test_append_commands_fields_vectors_and_insert_before_headers(tmp_path):
    workbook=load_workbook(PROFILE_PATH)
    for index in range(35):
        command(workbook['02_指令'],60+index,f'EXTRA{index}',f'{0x70+index:02X}')
        field(workbook['03_发送字段'],100+index,f'EXTRA{index}')
    vector=list(next(workbook['06_测试向量'].iter_rows(min_row=6,max_row=6,values_only=True)))
    for index in range(25):
        vector[1]=f'VECTOR{index}'
        row(workbook['06_测试向量'],40+index,vector)
    for name in ['01_基本信息','02_指令','03_发送字段','06_测试向量']:
        workbook[name].insert_rows(4,7)
    profile=save_load(workbook,tmp_path)
    assert len(profile.vector_results)==30
    for index in range(35):
        assert profile.encode(f'EXTRA{index}',0)[6]==77


def test_enum_and_bit_sections_expand_with_no_row_limit(tmp_path):
    workbook=load_workbook(PROFILE_PATH)
    sheet=workbook['05_枚举位域']
    for merged in list(sheet.merged_cells.ranges): sheet.unmerge_cells(str(merged))
    marker=next(r for r in range(1,sheet.max_row+1) if sheet.cell(r,1).value=='B. 位域布局（可选）')
    sheet.insert_rows(marker,60)
    for index in range(50):
        row(sheet,marker+index,['Y',f'EXTRA_ENUM{index}','01','CHOICE','选择','新增枚举'])
    for index in range(25):
        row(sheet,200+index,['Y',f'EXTRA_BITS{index}','value',0,8,'CONSTANT',index,None,'新增位域'])
    profile=save_load(workbook,tmp_path)
    assert len([key for key in profile.enums if key.startswith('EXTRA_ENUM')])==50
    assert profile.compile_bit_layout('EXTRA_BITS24',{})==bytes([24])


@pytest.mark.parametrize('kind',['enum_width','enum_alias','receive_enum','receive_duplicate','int_width','bit_overlap','bad_flag','bad_vector_after_limit'])
def test_invalid_extensions_fail_during_import_with_location(tmp_path,kind):
    workbook=load_workbook(PROFILE_PATH)
    if kind=='enum_width': workbook['05_枚举位域']['C13']='01 02'
    if kind=='enum_alias': workbook['05_枚举位域']['D13']='FRAME_SWITCH'
    if kind=='receive_enum':
        sheet=workbook['04_接收解析']
        index=next(r for r in range(1,sheet.max_row+1) if sheet.cell(r,4).value=='init_state')
        sheet.cell(index,12).value='UNKNOWN_ENUM'
    if kind=='receive_duplicate':
        sheet=workbook['04_接收解析']
        index=next(r for r in range(1,sheet.max_row+1) if sheet.cell(r,4).value=='init_state')
        row(sheet,150,[sheet.cell(index,c).value for c in range(1,14)])
    if kind=='int_width':
        command(workbook['02_指令'],60)
        field(workbook['03_发送字段'],100,kind='uint16',length=1,default=256)
    if kind=='bit_overlap': row(workbook['05_枚举位域'],100,['Y','CHIP_DIRECT_WRITE_WORD','overlap',0,8,'CONSTANT',1,None,'重叠'])
    if kind=='bad_flag': workbook['02_指令']['A60']='YES'
    if kind=='bad_vector_after_limit':
        values=[workbook['06_测试向量'].cell(6,c).value for c in range(1,11)]
        values[1]='BAD';values[7]='00 '*22
        row(workbook['06_测试向量'],100,values)
    with pytest.raises(ServiceError) as caught: save_load(workbook,tmp_path)
    assert caught.value.target


def test_manual_profile_and_generic_bitfield_binding(tmp_path):
    workbook=load_workbook(PROFILE_PATH)
    for sheet_name in ['02_指令','03_发送字段','06_测试向量']:
        sheet=workbook[sheet_name]
        for r in range(6,sheet.max_row+1):
            if sheet.cell(r,1).value=='Y': sheet.cell(r,1).value='N'
    command(workbook['02_指令'],60)
    field(workbook['03_发送字段'],100,kind='bitfield',length=1,default=None,layout='CUSTOM_BITS')
    row(workbook['05_枚举位域'],100,['Y','CUSTOM_BITS','upper',0,4,'MAPPING','level',None,'上4位'])
    row(workbook['05_枚举位域'],101,['Y','CUSTOM_BITS','lower',4,4,'CONSTANT',3,None,'下4位'])
    profile=save_load(workbook,tmp_path)
    assert profile.encode('EXTRA',2,{'extra.level':5})[6]==0x53
    summary=profile.summary()['commands'][0]['fields']
    assert summary[0]['key']=='extra.level'
    assert summary[0]['maximum']==15
    frame=profile.encode('EXTRA',2,{'extra.level':5})
    row(workbook['06_测试向量'],100,['Y','CUSTOM_VECTOR','TX','EXTRA','通用位域',2,'extra.level=5',frame.hex(),None,None])
    assert save_load(workbook,tmp_path).vector_results[0]['passed']


def test_rx_vectors_are_executed(tmp_path):
    workbook=load_workbook(PROFILE_PATH)
    profile=ProfileLoader().load(str(PROFILE_PATH))
    frame=profile.encode('QUERY_INIT_STATE',0,{})
    row(workbook['06_测试向量'],100,['Y','RX_INIT','RX','QUERY_INIT_STATE','接收枚举校验',0,'init_state=INITIALIZED',frame.hex(),None,None])
    profile=save_load(workbook,tmp_path)
    assert profile.vector_results[-1]['direction']=='RX'
    workbook['06_测试向量']['G100']='init_state=FAILED'
    with pytest.raises(ServiceError,match='RX'): save_load(workbook,tmp_path)


def test_all_decoded_fields_are_visible():
    profile=ProfileLoader().load(str(PROFILE_PATH))
    rule=profile.response_rules[0x09]
    original=profile.response_fields[rule.rule_id][0]
    profile.response_fields[rule.rule_id]=[replace(original,key=f'field{index}') for index in range(12)]
    registry=AssetRegistry()
    registry.add_profile(profile)
    assert len(registry.decode_matching_frame(profile.encode('QUERY_INIT_STATE',0,{}))['fields'])==12


def test_templates_use_expandable_tables_and_correct_validation_rules():
    for path in PROFILE_PATH.parent.glob('*V1.0.xlsx'):
        workbook=load_workbook(path)
        assert workbook.defined_names['ProfileEnums'].attr_text=='ProfileEnumRows[MappingId]'
        assert 'ProfileEnumRows' in workbook['05_枚举位域'].tables
        assert 'ProfileSendFields' in workbook['03_发送字段'].tables
        rules=workbook['06_测试向量'].data_validations.dataValidation
        assert any('C6' in str(rule.sqref) and rule.formula1=='"TX,RX"' for rule in rules)
        rules=workbook['03_发送字段'].data_validations.dataValidation
        assert not any('I6' in rule.sqref for rule in rules)
        assert any('O1048576' in rule.sqref and rule.formula1=='ProfileEnums' for rule in rules)
